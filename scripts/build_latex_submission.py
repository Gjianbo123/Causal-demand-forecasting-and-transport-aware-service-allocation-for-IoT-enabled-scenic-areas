"""Convert the final audited Word manuscript to a standalone Springer LaTeX source.

This is a format conversion: paragraphs and table values come from the final
DOCX; native math comes from the exact equation definitions used to build it.
No experiments, result selection, or statistical calculations are performed.
"""
from pathlib import Path
import ast
import hashlib
import json
import re
import shutil
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]
REV=ROOT/'output/revision'
OUT=REV/'Discover_IoT_LaTeX'
SOURCE=REV/'Discover_IoT_Revised_Manuscript.docx'
REFERENCES=json.loads((REV/'reference_order.json').read_text(encoding='utf-8'))
REF_KEYS={i:r['id'] for i,r in enumerate(REFERENCES,1)}


def tex(text, citations=True):
    protected={}
    def protect(code):
        key=f'LATEXPLACEHOLDER{len(protected)}END'
        protected[key]=code
        return key
    if citations:
        text=re.sub(r'\[(\d+(?:,\s*\d+)*)\]',lambda m:protect(r'\citep{'+','.join(REF_KEYS[int(x)] for x in m[1].split(','))+'}'),text)
        text=re.sub(r'\bFigure (\d+)\b',lambda m:protect(r'Figure~\ref{fig:'+m[1]+'}'),text)
        text=re.sub(r'\bTable (\d+)\b',lambda m:protect(r'Table~\ref{tab:'+m[1]+'}'),text)
    for old,new in {
        'ρ = πθ(u|s)/πold(u|s)':r'\(\rho=\pi_\theta(u\mid s)/\pi_{\mathrm{old}}(u\mid s)\)',
        'KL = mean[(ρ−1)−logρ]':r'\(\mathrm{KL}=\operatorname{mean}[(\rho-1)-\log\rho]\)',
        'h ∈ {1,2,4,6}':r'\(h\in\{1,2,4,6\}\)',
        '10⁻⁴':r'\(10^{-4}\)',
        'I⁺':r'\(I^+\)',
    }.items():
        if old in text: text=text.replace(old,protect(new))
    escaped={'\\':r'\textbackslash{}','{':r'\{','}':r'\}','$':r'\$','&':r'\&','#':r'\#','_':r'\_','%':r'\%','~':r'\textasciitilde{}','^':r'\textasciicircum{}',
        'Δ':r'\(\Delta\)','η':r'\(\eta\)','γ':r'\(\gamma\)','λ':r'\(\lambda\)','ρ':r'\(\rho\)','π':r'\(\pi\)','θ':r'\(\theta\)',
        'ℓ':r'\(\ell\)','≤':r'\(\le\)','∈':r'\(\in\)','×':r'\(\times\)','±':r'\(\pm\)',
        '−':'-','–':'--','—':'---','’':"'",'“':'``','”':"''",'é':r"\'{e}",'ñ':r'\~{n}','ò':r'\`{o}','ó':r"\'{o}",'ô':r'\^{o}','ö':r'\"{o}','ü':r'\"{u}','ć':r"\'{c}",'č':r'\v{c}'}
    result=''.join(escaped.get(c,c) for c in text)
    for key,value in protected.items(): result=result.replace(key,value)
    unknown=sorted(set(c for c in result if ord(c)>127))
    if unknown: raise ValueError(f'Unmapped Unicode: {unknown!r}')
    return result


def equations():
    tree=ast.parse((ROOT/'scripts/build_submission.py').read_text(encoding='utf-8'))
    eq=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='EQ' for t in n.targets))
    # Line breaks preserve the exact mathematics while fitting the 12 pt page.
    eq[1]=r'\begin{aligned}\lambda_{d,t}&=e_{d,t}+0.68\lambda_{d,t-1}P,\\ X_{d,t,i}&\sim\operatorname{Poisson}(\lambda_{d,t,i}\eta_{d,t,i}).\end{aligned}'
    eq[2]=r'\begin{aligned}R_{d,t,k,i}&\sim\operatorname{Binomial}(X_{d,t,i},p_k),\\p&=(0.75,0.55,0.40,0.30).\end{aligned}'
    eq[6]=r'\begin{aligned}\beta_{k,i,h}&=\operatorname{softmax}_{h}\left(\frac{q_k^{\mathsf T}\tanh(We_{i,h})}{\sqrt{8}}\right),\\f_{k,i}&=\sum_h\beta_{k,i,h}e_{i,h}.\end{aligned}'
    return eq


def bibliography():
    items=[]
    for r in REFERENCES:
        venue=r['venue']
        is_preprint='preprint' in r.get('publication_type','').lower() or 'arxiv' in venue.lower()
        conference=any(w in venue.lower() for w in ('conference','proceedings','symposium'))
        kind='misc' if is_preprint else ('inproceedings' if conference else 'article')
        protected_names={
            'Edwin Hamel-De le Court':r'{Hamel-De le Court}, Edwin',
            'Santiago Ontañón':r"{Onta\~{n}\'{o}n}, Santiago",
        }
        names=[protected_names.get(a,tex(a,False)) for a in r['authors']]
        fields={'author':' and '.join(names),
                'title':'{'+tex(r['title'],False)+'}', 'year':str(r['year'])}
        fields['howpublished' if is_preprint else ('booktitle' if conference else 'journal')]=tex(venue,False)
        if r.get('volume'): fields['volume']=str(r['volume'])
        if r.get('pages_or_article'): fields['pages']=tex(str(r['pages_or_article']),False)
        if r.get('doi'): fields['doi']=r['doi']
        elif r.get('url'): fields['url']=r['url']
        items.append('@'+kind+'{'+r['id']+',\n'+',\n'.join('  '+k+' = {'+v+'}' for k,v in fields.items())+'\n}')
    (OUT/'references.bib').write_text('\n\n'.join(items)+'\n',encoding='utf-8')


def table_source(table,number,caption):
    rows=[[cell.text for cell in row.cells] for row in table.rows]
    columns='@{}L'+('Y'*(len(rows[0])-1))+'@{}'
    def cell(value): return tex(value).replace('\n',r'\newline ')
    lines=[r'\begin{table}[!htbp]',r'\caption{'+tex(caption)+r'}\label{tab:'+str(number)+'}',
           r'\centering',r'\normalsize',r'\setlength{\tabcolsep}{4pt}',r'\renewcommand{\arraystretch}{1.18}',
           r'\begin{tabularx}{\textwidth}{'+columns+'}',r'\toprule',
           ' & '.join(r'\textbf{'+cell(x)+'}' for x in rows[0])+r' \\',r'\midrule']
    for row in rows[1:]: lines.append(' & '.join(cell(x) for x in row)+r' \\')
    lines += [r'\bottomrule',r'\end{tabularx}',r'\end{table}']
    return '\n'.join(lines)


PREAMBLE=r'''% Springer Nature template. Compile main.tex with pdfLaTeX + BibTeX.
% Numerical results and scientific content are unchanged from the audited DOCX.
\documentclass[pdflatex,sn-mathphys-num]{sn-jnl}
\usepackage{xcolor}
\usepackage{graphicx}
\usepackage{amsmath,amssymb}
\usepackage{booktabs,tabularx,array}
\usepackage{placeins}
\usepackage{microtype}
\usepackage{xurl}
\urlstyle{same}
\hypersetup{hidelinks}
\geometry{reset,a4paper,left=22.45mm,right=22.45mm,top=25mm,bottom=25mm}
\newcolumntype{L}{>{\raggedright\arraybackslash}X}
\newcolumntype{Y}{>{\centering\arraybackslash}X}
\graphicspath{{figures/}}
\raggedbottom
\setlength{\emergencystretch}{2em}
% The journal requires at least 12 pt. This is applied explicitly because the
% class does not necessarily implement standard article-class size options.
\makeatletter
\renewcommand\normalsize{\@setfontsize\normalsize{12}{15}%
  \abovedisplayskip 10pt plus 2pt minus 5pt
  \belowdisplayskip \abovedisplayskip
  \abovedisplayshortskip 0pt plus 3pt
  \belowdisplayshortskip 6pt plus 3pt minus 3pt
  \let\@listi\@listI}
\def\abstractfont{\reset@font\fontsize{12}{15}\selectfont\leftskip=0pt\rightskip=0pt\parfillskip=0pt plus 1fil}
\def\abstractheadfont{\reset@font\fontsize{12}{15}\bfseries\selectfont\titraggedcenter}
\def\abstractsubheadfont{\reset@font\fontsize{12}{15}\bfseries\selectfont}
\def\keywordfont{\reset@font\fontsize{12}{15}\selectfont\leftskip=0pt\rightskip=0pt plus .5fil}
\def\addressfont{\reset@font\fontsize{12}{15}\selectfont\titraggedcenter}
\def\figurecaptionfont{\reset@font\fontsize{12}{15}\selectfont}
\def\tablecaptionfont{\reset@font\fontsize{12}{15}\selectfont}
\def\tablebodyfont{\reset@font\fontsize{12}{15}\selectfont}
\def\tablecolheadfont{\reset@font\fontsize{12}{15}\bfseries\selectfont}
\def\tablefootnotefont{\reset@font\fontsize{12}{15}\selectfont}
\def\subsubsectionfont{\reset@font\fontsize{12}{15}\bfseries\selectfont}
\def\bibfont{\reset@font\normalsize\selectfont}
\makeatother
\AtBeginDocument{\normalsize}
\begin{document}
'''


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'figures').mkdir(exist_ok=True)
    template=ROOT/'tmp/latex_template/sn-article-template'
    shutil.copy2(template/'sn-jnl.cls',OUT/'sn-jnl.cls')
    shutil.copy2(template/'bst/sn-mathphys-num.bst',OUT/'sn-mathphys-num.bst')
    doc=Document(SOURCE)
    paragraphs=doc.paragraphs
    title=paragraphs[0].text
    abstract=next(paragraphs[i+1].text for i,p in enumerate(paragraphs) if p.text=='Abstract')
    keywords=next(p.text.removeprefix('Keywords: ') for p in paragraphs if p.text.startswith('Keywords:'))
    parts=[PREAMBLE,r'\title[Forecasts and IoT service allocation]{'+tex(title)+'}',
           r'\author*[1]{\fnm{Jianbo} \sur{Guo}}\email{hsng5987123@163.com}',
           r'\affil*[1]{\orgname{Yongcheng Vocational College}, \orgaddress{\city{Yongcheng}, \state{Henan}, \postcode{476600}, \country{China}}}',
           r'\abstract{'+tex(abstract)+'}',r'\keywords{'+tex(keywords)+'}',r'\maketitle']
    eq=equations(); ne=0; nf=0; nt=0; in_body=False; in_list=False; pending_caption=None
    converted=[]
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}tbl'):
            nt+=1
            assert pending_caption and pending_caption[0]==nt
            parts.append(table_source(Table(child,doc),nt,pending_caption[1])); pending_caption=None
            continue
        if not child.tag.endswith('}p'): continue
        p=Paragraph(child,doc); text=p.text
        if text=='1 Introduction': in_body=True
        if not in_body: continue
        if text=='References': break
        if child.xpath('.//m:oMath'):
            ne+=1
            parts.append(r'\begin{equation}\label{eq:'+str(ne)+'}\n'+eq[ne]+'\n'+r'\end{equation}')
            continue
        if child.xpath('.//w:drawing'): continue
        if not text.strip(): continue
        is_list=p.style.name=='List Bullet'
        if in_list and not is_list: parts.append(r'\end{itemize}'); in_list=False
        if is_list and not in_list: parts.append(r'\begin{itemize}'); in_list=True
        if p.style.name.startswith('Heading'):
            level=int(p.style.name[-1]); label=re.sub(r'^\d+(?:\.\d+)*\s+','',text)
            # Allow the architecture/demand floats to share pages with methods.
            # Other section boundaries keep result figures near their discussion.
            if level==1 and label!='Forecasting and allocation methods': parts.append(r'\FloatBarrier')
            command='section' if level==1 else 'subsection'
            if label=='Declarations': command+='*'
            parts.append('\\'+command+'{'+tex(label)+'}')
        elif p.style.name=='Caption':
            m=re.match(r'Fig\. (\d+) (.*)',text)
            if m:
                nf+=1; assert nf==int(m[1]); filename=next((REV/'figures').glob(f'Fig{nf}_*.pdf'))
                shutil.copy2(filename,OUT/'figures'/filename.name)
                parts.append('\n'.join([r'\begin{figure}[!htbp]',r'\centering',
                    r'\includegraphics[width=\textwidth]{'+filename.name+'}',
                    r'\caption{'+tex(m[2])+r'}\label{fig:'+str(nf)+'}',r'\end{figure}']))
            else:
                m=re.match(r'Table (\d+) (.*)',text); assert m,text
                pending_caption=(int(m[1]),m[2])
        else:
            parts.append((r'\item ' if is_list else '')+tex(text))
        converted.append(text)
    if in_list: parts.append(r'\end{itemize}')
    assert (ne,nf,nt)==(9,5,4),(ne,nf,nt)
    parts += [r'\FloatBarrier',r'\bibliography{references}',r'\end{document}']
    (OUT/'main.tex').write_text('\n\n'.join(parts)+'\n',encoding='utf-8')
    bibliography()
    cover=Document(REV/'Discover_IoT_Cover_Letter.docx')
    cover_parts=[r'\documentclass[12pt,a4paper]{article}',r'\usepackage[margin=25mm]{geometry}',r'\usepackage{hyperref}',r'\setlength{\parindent}{0pt}',r'\setlength{\parskip}{8pt}',r'\setlength{\emergencystretch}{3em}',r'\raggedright',r'\begin{document}']
    for p in cover.paragraphs:
        if not p.text: continue
        value=tex(p.text,False).replace('\n',r'\\ ')
        cover_parts.append(r'{\large\bfseries '+value+'}' if p.style.name=='Title' else value)
    cover_parts.append(r'\end{document}')
    (OUT/'cover_letter.tex').write_text('\n\n'.join(cover_parts)+'\n',encoding='utf-8')
    (OUT/'conversion_audit.json').write_text(json.dumps({'source_docx':SOURCE.name,'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
       'paragraphs_converted':len(converted),'native_equations':ne,'figures':nf,'tables':nt,'references':len(REFERENCES),'results_recomputed':False},indent=2),encoding='utf-8')
    print(json.dumps({'output':str(OUT),'equations':ne,'figures':nf,'tables':nt,'references':len(REFERENCES)}))


if __name__=='__main__': main()
