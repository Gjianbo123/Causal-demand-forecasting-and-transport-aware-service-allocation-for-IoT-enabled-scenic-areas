"""Build the journal manuscript from locked experiment outputs, never old tables.

Run using the Codex bundled Python runtime (python-docx/lxml), not experiment Python.
"""
from pathlib import Path
import csv
import json
import re
import sys
from copy import deepcopy
from statistics import mean

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/revision'
RUN=ROOT/'runs/scenic_rebuild_v2'
sys.path.insert(0,str(ROOT/'tmp/doc_deps'))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
import latex2mathml.converter

TITLE='Evaluating the use of forecasts for service allocation under imperfect IoT sensing in a scenic area simulation'
LABELS={'F0':'Reactive PPO','F1':'Concatenation','F2':'Uniform fusion','F3':'Shared attention','F4':'Resource attention'}
EQ={
1:r'\lambda_{d,t}=e_{d,t}+0.68\lambda_{d,t-1}P,\quad X_{d,t,i}\sim\operatorname{Poisson}(\lambda_{d,t,i}\eta_{d,t,i})',
2:r'R_{d,t,k,i}\sim\operatorname{Binomial}(X_{d,t,i},p_k),\quad p=(0.75,0.55,0.40,0.30)',
3:r'Q_{k,i,t+1}=\max\{Q_{k,i,t}+R_{k,i,t+1}-2I^{+}_{k,i,t},0\}',
4:r'\tau_{k,i,j}=\left\lceil\frac{D_{i,j}/v_k+b_k}{10}\right\rceil,\quad i\ne j',
5:r'a_{i,j,t}=\operatorname{softmax}_{j}\left(\operatorname{LeakyReLU}\left(qz_{i,t}+kz_{j,t}+w\log(1+\ell_{i,j})+b\right)\right)',
6:r'\beta_{k,i,h}=\operatorname{softmax}_{h}\left(\frac{q_k^{\mathsf T}\tanh(We_{i,h})}{\sqrt{8}}\right),\quad f_{k,i}=\sum_h\beta_{k,i,h}e_{i,h}',
7:r'B_{k,i}\le A_{k,i}\le H_{k,i},\quad\sum_i A_{k,i}=N_k',
8:r'r=-\left[\frac{1}{25\cdot16}\sum_{i=1}^{16}O_i+0.025M+\frac{0.02}{16}\sum_{i=1}^{16}\max(O_i-50,0)\right]',
9:r'W_{\mathrm{restricted}}=\frac{10\sum_{t=0}^{71}\sum_{k,i}Q_{k,i,t}}{\sum_{t=1}^{72}\sum_{k,i}R_{k,i,t}}'
}
TRANSFORM=etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
refs=json.loads((OUT/'references_verified.json').read_text(encoding='utf-8'))['references']
refs.append(dict(number=26,id='pal2024',authors=['Shantanu Pal','Sara Khalifa','Dimity Miller','Volkan Dedeoglu','Ali Dorri','Gowri Ramachandran','Peyman Moghadam','Brano Kusy','Raja Jurdak'],title='Uncertainty propagation in the internet of things',year=2024,venue='Discover Internet of Things',volume='4',pages_or_article='32',doi='10.1007/s43926-024-00085-2',url='https://link.springer.com/article/10.1007/s43926-024-00085-2'))
reference_order=[]
all_prose=[]


def cite(text):
    def rep(match):
        nums=[int(n) for n in match.group(1).split(',')]
        updated=[]
        for n in nums:
            if n not in reference_order: reference_order.append(n)
            updated.append(str(reference_order.index(n)+1))
        return '['+', '.join(updated)+']'
    return re.sub(r'\[(\d+(?:,\s*\d+)*)\]',rep,text)


def setup():
    doc=Document()
    for style in doc.styles:
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.27),Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.85)
    sec.left_margin=sec.right_margin=Inches(.85)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption','List Bullet'):
        s=doc.styles[name]; s.font.name='Arial'; s.font.size=Pt(12); s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(17)
    doc.styles['Title'].font.bold=True
    for name in ('Heading 1','Heading 2','Heading 3'):
        doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(10)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.italic=False
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
    footer._p.append(fld)
    doc.core_properties.author='Jianbo Guo'
    doc.core_properties.title=TITLE
    return doc


def para(doc,text,style=None):
    text=text.replace('not a reconstruction of the unavailable original simulator or a calibrated digital twin',
                      'with an explicit generator and no real-site calibration')
    text=text.replace(', does not recover the earlier unavailable implementation, and ', ' and ')
    text=cite(text)
    all_prose.append(text)
    p=doc.add_paragraph(text,style)
    p.paragraph_format.widow_control=True
    return p


def heading(doc,text,level=1):
    all_prose.append(text)
    return doc.add_heading(text,level)


def equation(doc,n):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    mathml=latex2mathml.converter.convert(EQ[n])
    mathroot=etree.fromstring(mathml.encode())
    # Office's MathML transform treats multi-character <mo> words as delimiter
    # separators and can silently truncate them. Named operators are identifiers.
    for node in mathroot.iter('{http://www.w3.org/1998/Math/MathML}mo'):
        if node.text and len(node.text)>1 and node.text.isalpha():
            node.tag='{http://www.w3.org/1998/Math/MathML}mi'
            node.set('mathvariant','normal')
    omml=TRANSFORM(mathroot)
    p._p.append(deepcopy(omml.getroot()))
    p.add_run(f'  ({n})')
    p.paragraph_format.space_before=Pt(4)
    p.paragraph_format.space_after=Pt(6)


def table(doc,caption,headers,rows,widths):
    p=para(doc,caption,'Caption'); p.paragraph_format.keep_with_next=True
    tab=doc.add_table(rows=1,cols=len(headers)); tab.alignment=WD_TABLE_ALIGNMENT.CENTER
    tab.autofit=False
    for col,w in zip(tab.columns,widths): col.width=Inches(w)
    for i,h in enumerate(headers): tab.rows[0].cells[i].text=h
    for row in rows:
        cells=tab.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value)
    for ri,row in enumerate(tab.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'70'); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9EDF0'); tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.line_spacing=1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size=Pt(12); run.bold=(ri==0)
        cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return tab


def figure(doc,file,caption,width=6.45):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(OUT/'figures'/file),width=Inches(width))
    para(doc,caption,'Caption')


def fmt(x): return f'{x:.2f}'
def msd(v): return f"{v['mean']:.2f} ± {v['sd']:.2f}"
def pval(p): return '<0.001' if p<.001 else f'{p:.3f}'


def main():
    r=json.loads((OUT/'manuscript_results.json').read_text(encoding='utf-8'))
    assert r['total_policy_episodes']==7830
    assert r['total_baseline_episodes']==729
    assert r['executed_infeasible_epochs']==0
    assert all(b['executable_infeasible_pct']==0 for b in r['baselines'].values())
    forecasts=list(csv.DictReader((RUN/'assets/forecast/metrics.csv').open(encoding='utf-8')))
    def fm(model,h,metric='mae'):
        return mean(float(x[metric]) for x in forecasts if x['split']=='test' and x['model']==model and int(x['horizon_steps'])==h)
    models=['persistence','historical_average','ridge','lstm','gru','edge_stgru']
    display=['Persistence','Historical mean','Ridge','LSTM','GRU','Edge STGRU']
    policies=r['policies']; base=r['baselines']
    contrast=next(c for c in r['contrasts'] if c['comparison']=='F4_minus_F1' and c['condition']=='nominal')
    delta=contrast['mean']; primary_p=contrast['p_value']
    edge=fm('edge_stgru',6); ridge=fm('ridge',6)
    f4=policies['F4']['restricted_mean_wait_min']['mean']; f1=policies['F1']['restricted_mean_wait_min']['mean']
    best_policy=min(policies,key=lambda m:policies[m]['restricted_mean_wait_min']['mean'])
    best_baseline=min(base,key=lambda m:base[m]['restricted_mean_wait_min'])
    bestbase=base[best_baseline]['restricted_mean_wait_min']
    evidence=('did not demonstrate a nominal waiting-time advantage over concatenation' if primary_p>=.05 or delta>=0
              else 'reduced nominal accrued waiting relative to concatenation')
    abstract=("IoT service allocation must convert uncertain observations into executable resource movements. "
      "This study evaluates forecast fusion in a reproducible scenic-area simulation with 24 nodes, four resource classes, "
      "travel delays and finite operating days. Six forecasting methods predict inflow 10, 20, 40 and 60 minutes ahead. "
      "Five proximal policy optimization variants compare absent forecasts, concatenation, uniform fusion, shared attention "
      "and resource-specific attention. All five variants are evaluated under nominal sensing; F0, F1 and F4 also face eight sensing faults, "
      "using ten policy seeds and 27 held-out days. "
      "The primary endpoint includes elapsed waiting of requests unserved at closing. "
      f"At 60 minutes, edge-STGRU and ridge regression obtain MAEs of {edge:.2f} and {ridge:.2f} visit events. "
      f"Resource-specific attention gives {f4:.2f} minutes of accrued waiting per request versus {f1:.2f} for concatenation "
      f"(paired difference {delta:+.2f} minutes; 95% confidence interval {contrast['ci_low']:.2f} to {contrast['ci_high']:.2f}). "
      f"The {best_baseline} heuristic obtains {bestbase:.2f} minutes. "
      f"All {r['policy_decisions']:,} learned-policy actions pass the implemented resource-invariant checks. "
      "The findings distinguish action feasibility from service quality and show why forecast fusion requires direct operational validation. "
      "The experiment is fully synthetic and does not establish field effectiveness.")
    assert len(abstract.split())<250,len(abstract.split())
    doc=setup()
    doc.add_paragraph(TITLE,'Title')
    para(doc,'Jianbo Guo')
    para(doc,'Yongcheng Vocational College, Yongcheng, Henan 476600, China')
    para(doc,'Corresponding author: hsng5987123@163.com')
    heading(doc,'Abstract')
    para(doc,abstract)
    para(doc,'Keywords: Internet of Things; imperfect sensing; spatiotemporal forecasting; resource allocation; reinforcement learning; synthetic benchmark')
    literature=(OUT/'literature_draft.md').read_text(encoding='utf-8').split('## Manuscript text')[1].split('## Integration notes')[0].strip().split('\n\n')
    heading(doc,'1 Introduction')
    para(doc,literature[0])
    para(doc,'Uncertainty can propagate across sensing, communication, processing and decision making in an IoT system [26]. '
         'A useful evaluation must therefore preserve when information becomes available and measure the resulting decisions, '
         'rather than identifying operational benefit solely from a lower prediction error.')
    research_questions=literature[-1].replace('Its contribution is an explicit and inspectable connection between demand generation, observation availability, forecasts, executable resource movements and measured service outcomes. ','').replace('The benchmark is intended to support reproducible mechanism comparisons. ','')
    para(doc,research_questions)
    para(doc,'The study makes three empirical contributions: a disclosed service-request simulator with transport commitments and causal telemetry; '
         'a controlled comparison isolating five forecast-fusion choices; and an evaluation that includes terminal backlog, sensing faults, '
         'independent policy seeds and executable-action audits. The research contribution is this testable integration and its observed limits.')
    heading(doc,'2 Related work')
    for p in literature[1:-1]: para(doc,p)
    sections=(OUT/'methods_draft.md').read_text(encoding='utf-8').split('\n## ')[1:]
    for si,section in enumerate(sections,3):
        title,body=section.split('\n',1)
        heading(doc,f'{si} {title.strip()}')
        for p in body.strip().split('\n\n'):
            eq=re.search(r'\((\d+)\)\.$',p.strip()) or re.search(r'\((\d+)\)$',p.strip())
            if eq and int(eq.group(1)) in EQ and len(p)<350:
                equation(doc,int(eq.group(1)))
            elif p.startswith('- '):
                for item in p.splitlines(): para(doc,item.removeprefix('- '),'List Bullet')
            else:
                p=p.replace('rather than the earlier E-STGNN','with an explicitly specified architecture')
                p=p.replace('No held-out policy result selected these settings or checkpoints.',
                  'No held-out policy result selected these settings or checkpoints. The uncertainty estimates condition on this one synthetic dataset and corruption realization.')
                para(doc,p)
        if si==3:
            para(doc,'Figure 1 summarizes the implemented sensing-to-action pipeline and physical graph. Figure 2 shows the generated demand and chronological data split.')
            figure(doc,'Fig1_architecture.png','Fig. 1 Implemented sensing to action pipeline and 24-node physical network. '
                   'Service locations, hubs and gates have different roles. The exact inventory ledger is separate from corrupted service-count telemetry')
            figure(doc,'Fig2_demand.png','Fig. 2 Synthetic demand and chronological evaluation design. (a) Daily totals of service requests; '
                   '(b) mean node visit-event counts by operating interval. Neither quantity represents unique tourists. '
                   'The train, validation and test intervals are generated from the same disclosed process')
        if si==5:
            para(doc,'For each paired contrast, an unadjusted two-sided 95% Student t interval is computed across the ten trained-seed differences '
                 '(nine degrees of freedom). These pointwise intervals describe conditional uncertainty and are not adjusted for multiple comparisons. '
                 'The sign-enumeration test remains the prespecified significance test. An exploratory transport sensitivity analysis evaluates '
                 'the frozen F1 and F4 policies with travel-duration multipliers 0.5 and 1.5, preserving the training normalization and test requests.')
            para(doc,'The experiments ran on Windows with 15.73 GiB physical RAM, an Intel64 Family 6 Model 186 processor and an NVIDIA '
                 'GeForce RTX 4060 Laptop GPU. Software versions were Python 3.11.5, PyTorch 2.7.1 with CUDA 11.8, NumPy 1.26.4, '
                 'SciPy 1.11.4, pandas 2.2.3 and Matplotlib 3.10.7. The release records configuration, source and checkpoint hashes.')
    heading(doc,'6 Results')
    heading(doc,'6.1 Forecasting performance',2)
    rows=[[label]+[fmt(fm(model,h)) for h in (1,2,4,6)]+[fmt(fm(model,6,'rmse'))] for model,label in zip(models,display)]
    table(doc,'Table 1 Held-out forecasting errors in visit events. Neural entries average three trained seeds; all models use the same 1,809 test origins',
          ['Model','10 min\nMAE','20 min\nMAE','40 min\nMAE','60 min\nMAE','60 min\nRMSE'],rows,[1.5,1,1,1,1,1])
    para(doc,f'In Table 1, edge-STGRU obtains a 60-minute MAE of {edge:.2f}, compared with {fm("gru",6):.2f} for GRU and {ridge:.2f} for ridge regression. '
         f'Its 60-minute WAPE is {fm("edge_stgru",6,"wape_pct"):.2f}%, versus {fm("ridge",6,"wape_pct"):.2f}% for ridge. '
         'Figure 3 presents horizon-specific errors and the graph versus graph-free comparison. These results do not establish dominance '
         'over modern Transformer forecasters, which are discussed as related work but are not reimplemented in this experiment.')
    para(doc,'The generator’s linear propagation and smooth periodic forcing may help explain the competitive performance of ridge regression. '
         'The edge-STGRU seed-zero forecast is retained for every policy comparison because it was fixed before policy evaluation.')
    figure(doc,'Fig3_forecasting.png','Fig. 3 Forecasting comparison at all four horizons and paired graph versus graph-free errors. '
           'Neural summaries use three training seeds; the paired daily panel averages those seeds. Differences are Edge-STGRU minus GRU; '
           'negative values favor Edge-STGRU. The orange line is the mean across 27 test days. Error bars denote sample SD, not a confidence interval')
    heading(doc,'6.2 Allocation and fusion',2)
    prows=[]
    for m in LABELS:
        a=policies[m]
        prows.append([m,msd(a['restricted_mean_wait_min']),msd(a['unserved_pct']),msd(a['visitor_overload_pct']),msd(a['movement_cost'])])
    for m in ('static','reactive','lookahead'):
        a=base[m]; prows.append([m.capitalize(),fmt(a['restricted_mean_wait_min']),fmt(a['unserved_pct']),fmt(a['visitor_overload_pct']),fmt(a['movement_cost'])])
    table(doc,'Table 2 Operational results under nominal sensing. Learned methods show mean ± sample SD across ten seeds; deterministic heuristics show pooled results on the same 27 days',
          ['Method','Accrued wait\nmin/request','Unserved\n%','Backlog\nexceedance %','Movement\nresource km/day'],prows,[1.05,1.5,1.25,1.4,1.3])
    para(doc,f'In the nominal comparison (Table 2), resource-specific attention {evidence}. Its primary difference from F1 is {delta:+.2f} min/request '
         f'(95% t interval {contrast["ci_low"]:.2f} to {contrast["ci_high"]:.2f}; exact paired p {pval(primary_p)}). '
         f'The lowest mean among the five learned variants is {best_policy}, at '
         f'{policies[best_policy]["restricted_mean_wait_min"]["mean"]:.2f} min/request. '
         'Figure 4 shows every trained seed; Table 3 gives paired contrasts and intervals.')
    if bestbase<policies[best_policy]['restricted_mean_wait_min']['mean']:
        para(doc,f'The {best_baseline} heuristic has lower accrued waiting than every learned variant in this experiment. '
             'Accordingly, the evidence does not support replacing this heuristic with the tested PPO controllers. '
             'The observation emphasizes the service availability lost during resource transport.')
    else:
        para(doc,'The heuristic comparison provides context beyond differences among learned fusion modules. '
             'It does not identify a globally optimal controller: no mixed-integer MPC or exhaustive allocation optimum is claimed.')
    para(doc,'The primary endpoint and served-only waiting answer different questions. For static allocation, accrued waiting is '
         f'{base["static"]["restricted_mean_wait_min"]:.2f} min/request whereas completed-request mean waiting is '
         f'{base["static"]["mean_wait_min"]:.2f} min; {base["static"]["unserved_pct"]:.2f}% of requests remain unserved. '
         'The difference illustrates why closing backlog must accompany waiting-time claims. Complete served-request P95 values and per-seed outputs are in the reproducibility archive.')
    figure(doc,'Fig4_allocation.png','Fig. 4 Nominal allocation outcomes. Each point is one independently trained policy evaluated across the same 27 days; black segments denote means. '
           'F0 has no forecasts; F1 concatenates; F2 averages; F3 shares an attention query; F4 uses resource-specific queries')
    crows=[]
    for c in r['contrasts']:
        if c['condition']=='nominal':
            crows.append([c['comparison'].replace('_minus_',' minus '),fmt(c['mean']),f"{c['ci_low']:.2f} to {c['ci_high']:.2f}",pval(c['p_value']),pval(c['p_holm']) if c['p_holm'] is not None else 'Primary'])
    table(doc,'Table 3 Paired nominal fusion contrasts for accrued waiting. Negative differences favor F4. The two secondary contrasts use Holm correction',
          ['Contrast','Difference\nmin','Unadjusted 95% t interval\nmin','Exact p','Adjusted p'],crows,[1.6,1.1,1.7,1.05,1.05])
    heading(doc,'6.3 Sensing and travel sensitivity',2)
    para(doc,'Figure 5 applies faults only at evaluation time. The eight sensing comparisons retain the same physical demand trajectories, '
         'inventory rules and frozen model parameters. A fault need not monotonically increase waiting. One possible explanation is that '
         'stale observations suppress relocations; this mechanism has not been isolated experimentally. Any apparent improvement under degradation '
         'does not establish that reliable observations are undesirable.')
    figure(doc,'Fig5_sensing.png','Fig. 5 Accrued waiting under separately applied sensing faults. Curves are means and error bars are sample SD across ten training seeds. '
           'Packet loss is per whole-node packet, noise is scaled by training channel SD, and delay preserves original event times')
    sig=sum(c['p_holm'] is not None and c['p_holm']<.05 for c in r['contrasts'] if c['family']=='robustness_eight')
    para(doc,f'Of the eight prespecified F4-minus-F1 sensing contrasts, {sig} have Holm-adjusted p below 0.05. '
         'The complete effect estimates, directions and intervals are supplied with the experiment outputs; a significance count alone does '
         'not establish operational superiority or robustness to combined, correlated or adversarial faults.')
    if r['sensitivity']:
        srows=[]
        for scale in (.5,1.5):
            vals={m:[x['restricted_mean_wait_min'] for x in r['sensitivity'] if x['method']==m and x['travel_multiplier']==scale] for m in ('F1','F4')}
            srows.append([str(scale),fmt(mean(vals['F1'])),fmt(mean(vals['F4'])),fmt(mean(vals['F4'])-mean(vals['F1']))])
        table(doc,'Table 4 Accrued waiting with altered transport durations. Frozen policies are evaluated without retraining; values average ten seeds',
              ['Travel multiplier','F1 min/request','F4 min/request','F4 minus F1'],srows,[1.8,1.5,1.5,1.7])
        para(doc,'The interventions in Table 4 change transport duration while preserving the training ledger normalization. '
             'They evaluate sensitivity to a specified mismatch and do not demonstrate transfer to another network, resource budget or destination.')
    heading(doc,'6.4 Feasibility and computational cost',2)
    para(doc,f'All {r["policy_decisions"]:,} learned-policy actions pass the implemented resource-invariant checks; '
         'the three heuristics also pass these checks. Deterministic target repair, reserved destination capacity '
         'and inventory conservation assertions enforce the modeled constraints. This is not an independent field-safety certification and '
         'does not imply zero service backlog. Candidate repair frequency for F4 is '
         f'{policies["F4"]["repair_pct"]["mean"]:.2f}%, illustrating the distinction between proposed scores and executable targets.')
    para(doc,f'F4 policy inference averages {policies["F4"]["policy_latency_ms"]["mean"]:.2f} ms per decision, and the instrumented local '
         f'telemetry-to-dispatch computation averages {policies["F4"]["local_pipeline_latency_ms"]["mean"]:.2f} ms. '
         'These timings were gathered during concurrent CPU evaluations and are descriptive rather than isolated hardware benchmarks. '
         'They exclude physical sensing, data transmission, service simulation and actuation. These measurements do not establish '
         'end-to-end field latency or performance on a deployed edge device.')
    heading(doc,'7 Discussion and limitations')
    para(doc,'The experiments separate three propositions that should not be conflated: a model can predict reasonably, a controller can produce '
         'feasible movements, and the movements can improve service. Success at either of the first two does not establish the third. '
         'Resource-specific attention is a testable fusion mechanism, but its operational value depends on transport cost, allocation dynamics '
         'and optimization quality. The heuristic baselines show whether differences among PPO variants translate into gains over simpler allocation rules.')
    para(doc,'The transport model makes frequent redistribution costly in service availability. Every dispatched unit loses at least one '
         'service interval, and long trips reserve destination capacity. Reactive or forecast-informed target tracking can therefore increase '
         'waiting even when it follows demand more closely. A controller that explicitly values persistent assignments or plans transport '
         'commitments may improve this trade-off, but that hypothesis is not established by the current results.')
    para(doc,'The strongest limitation is external validity. The single graph and generator are not calibrated to a real scenic area; '
         'visit events do not form conserved tourist trajectories, and independent request classes are abstractions. The mean demand process '
         'contains a linear graph component, potentially favoring linear forecasting. One held-out block and one fault realization per day '
         'cannot represent the range of demand shifts or communication failures in deployed systems. Ten policy seeds quantify optimization '
         'variation conditional on this environment, not variability across real destinations.')
    para(doc,'Training budgets and architectures were common across fusion variants but were not exhaustively optimized. The separately trained '
         'forecaster was not selected for downstream reward, and the attention comparisons do not isolate every parameter-count difference. '
         'There is no claim of state-of-the-art forecasting, certified crowd safety, optimal scheduling or global equivalence to MPC. '
         'Future validation should use permissioned sensor traces, measured service and transport times, graph and demand shifts, '
         'strong model-based controllers and longer training-budget studies. Deployment would additionally require failure handling for '
         'resource-location and dispatch-confirmation errors, which the current trusted ledger excludes.')
    heading(doc,'8 Conclusion')
    para(doc,f'This study provides a reproducible synthetic evaluation of forecast fusion for IoT-informed service allocation. '
         f'The primary resource-attention versus concatenation difference is {delta:+.2f} min/request under nominal sensing, '
         f'and the {best_baseline} heuristic achieves {bestbase:.2f} min/request. '
         'The combined forecasting, terminal-backlog, fault and feasibility results support evaluating the complete sensing-to-action chain '
         'rather than inferring operational improvement from forecast accuracy alone. The evidence remains conditional on the disclosed '
         'simulation and motivates further model and field validation.')
    heading(doc,'Declarations')
    declarations=OUT/'author_declarations.json'
    author=json.loads(declarations.read_text()) if declarations.exists() else {}
    para(doc,'Funding: '+author.get('funding','Author confirmation required before submission.'))
    para(doc,'Competing interests: '+author.get('competing_interests','Author confirmation required before submission.'))
    para(doc,'Author contributions: '+author.get('contributions','The corresponding author must confirm the contribution statement before submission.'))
    para(doc,'Ethics and consent: This computational study uses wholly synthetic data and involves no human participants, identifiable personal data or animal subjects.')
    para(doc,'Data and code availability: The synthetic-data generator, configuration, model checkpoints, raw episode outputs, analysis scripts '
         'and figure sources are provided in the accompanying reproducibility archive. No real sensor dataset was used. A public archival '
         'repository identifier has not been assigned; access arrangements must be finalized before publication.')
    para(doc,'AI assistance: OpenAI Codex assisted with the new software implementation, experiment orchestration, analysis and manuscript preparation. '
         'The underlying configuration, code and numerical outputs are supplied for inspection. Final scientific review and approval remain the author’s responsibility.')
    heading(doc,'References')
    for i,n in enumerate(reference_order,1):
        ref=next(x for x in refs if x['number']==n)
        authors='; '.join(ref['authors'])
        details=f"{authors}. {ref['title']}. {ref['venue']}. {ref['year']}"
        if ref.get('volume'): details+=';'+str(ref['volume'])
        if ref.get('pages_or_article'): details+=':'+ref['pages_or_article']
        link='https://doi.org/'+ref['doi'] if ref.get('doi') else ref['url']
        p=doc.add_paragraph(f'{i}. {details}. {link}')
        p.paragraph_format.left_indent=Inches(.24)
        p.paragraph_format.first_line_indent=Inches(-.24)
        p.paragraph_format.space_after=Pt(6)
    target=OUT/'Discover_IoT_Revised_Manuscript.docx'
    doc.save(target)
    (OUT/'manuscript_text.txt').write_text('\n\n'.join(all_prose),encoding='utf-8')
    (OUT/'reference_order.json').write_text(json.dumps([next(x for x in refs if x['number']==n) for n in reference_order],indent=2),encoding='utf-8')
    cover=setup(); cover.core_properties.title='Submission to Discover Internet of Things'
    cover.add_paragraph('Submission to Discover Internet of Things','Title')
    para(cover,'Dear Editors,')
    para(cover,f'Please consider the Research article entitled “{TITLE}” by Jianbo Guo.')
    para(cover,'The manuscript examines how incomplete and delayed IoT observations affect forecast-informed resource allocation. '
         'It connects an explicit synthetic demand generator, causal telemetry, four-horizon prediction, learned allocation preferences '
         'and deterministic resource-feasibility repair. The evaluation includes six forecasters, five PPO fusion variants, ten policy '
         'seeds, nine sensing conditions and transparent heuristic comparators.')
    para(cover,f'The principal comparison gives an F4-minus-F1 accrued-wait difference of {delta:+.2f} min/request under nominal sensing. '
         'The paper retains adverse and null findings, distinguishes action feasibility from service quality, and supplies implementation '
         'and raw outputs. Its fit with the journal lies in the analysis of the IoT sensing-to-action pipeline and performance modeling of '
         'an application, with clear limits on real-world inference.')
    para(cover,'The previous version of this work was rejected by Journal of Applied Science and Engineering. Because its original source '
         'implementation was unavailable, the present version replaces the numerical evidence with newly implemented and executed '
         'synthetic experiments. It does not claim recovery or independent reproduction of the earlier results.')
    para(cover,'OpenAI Codex assisted with software implementation and preparation; the manuscript contains an AI-assistance disclosure. '
         'The corresponding author will complete the required authorship, funding, competing-interest and exclusive-submission declarations '
         'and approve the manuscript before it is submitted.')
    para(cover,'Sincerely,\nJianbo Guo\nYongcheng Vocational College\nhsng5987123@163.com')
    cover.save(OUT/'Discover_IoT_Cover_Letter.docx')
    print(json.dumps(dict(manuscript=str(target),abstract_words=len(abstract.split()),references=len(reference_order),equations=len(doc._element.xpath('//m:oMath')),paragraphs=len(doc.paragraphs),tables=len(doc.tables))))


if __name__=='__main__': main()
